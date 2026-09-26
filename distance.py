"""Local distance helpers. No server is needed; import and call the functions.

Only read_projects_from_excel requires openpyxl. Calculations use the standard
library. All coordinates are latitude/longitude in decimal degrees.
"""

from datetime import date, datetime
from itertools import combinations
from math import asin, cos, isfinite, radians, sin, sqrt
from pathlib import Path
from typing import NotRequired, TypedDict


class Project(TypedDict):
    project_id: str
    utility: str
    lat_center: float
    lon_center: float
    project_name: NotRequired[str]
    in_service_date: NotRequired[str | date | None]


class DistanceResult(TypedDict):
    distance_km: float
    distance_mi: float
    is_overlap: bool


class ProjectPair(DistanceResult):
    project_id_a: str
    project_id_b: str
    utility_a: str
    utility_b: str
    project_name_a: str
    project_name_b: str
    time_gap_days: int | None


def _number(value: object, name: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(f"{name} must be a finite number") from exc
    if isinstance(value, bool) or not isfinite(result):
        raise ValueError(f"{name} must be a finite number")
    return result


def _point(lat: object, lon: object) -> tuple[float, float]:
    lat, lon = _number(lat, "latitude"), _number(lon, "longitude")
    if not -90 <= lat <= 90 or not -180 <= lon <= 180:
        raise ValueError("Latitude must be within [-90, 90] and longitude within [-180, 180]")
    return lat, lon


def _threshold(value: float) -> float:
    value = _number(value, "threshold_mi")
    if value <= 0:
        raise ValueError("threshold_mi must be positive")
    return value


def calculate_distance(
    lat_a: float, lon_a: float, lat_b: float, lon_b: float,
    threshold_mi: float = 25.0,
) -> DistanceResult:
    """Calculate the Haversine distance between two points.

    Input: four decimal-degree coordinates and a positive cutoff in miles.
    Output: distance_km, distance_mi, and is_overlap (distance_mi < cutoff).
    Distances are unrounded, using a mean Earth radius of 6371.0088 km.
    This is geographic distance, not driving distance or transmission-line length.
    Raises ValueError for missing, non-finite, or out-of-range coordinates.
    """
    lat_a, lon_a = _point(lat_a, lon_a)
    lat_b, lon_b = _point(lat_b, lon_b)
    threshold_mi = _threshold(threshold_mi)
    dlat, dlon = radians(lat_b - lat_a), radians(lon_b - lon_a)
    h = sin(dlat / 2) ** 2 + cos(radians(lat_a)) * cos(radians(lat_b)) * sin(dlon / 2) ** 2
    km = 2 * 6371.0088 * asin(sqrt(min(1.0, max(0.0, h))))
    miles = km / 1.609344
    return {"distance_km": km, "distance_mi": miles, "is_overlap": miles < threshold_mi}


def _date(value: str | date | None) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError("in_service_date must be YYYY-MM-DD, M/D/YYYY, a date, or None")


def compare_projects(
    projects: list[Project], threshold_mi: float = 25.0,
    include_non_overlaps: bool = False,
) -> list[ProjectPair]:
    """Compare centers of projects belonging to different utilities.

    Input: dictionaries with project_id, utility, lat_center, and lon_center.
    Optional fields: project_name and in_service_date (ISO/US date or date object).
    Project IDs must be unique; utility names are case-sensitive.
    Output: unique pairs sorted by distance, with IDs, names, utilities,
    distance_km, distance_mi, is_overlap, and time_gap_days. The date gap is
    absolute, or None if either date is missing; it is not construction overlap.
    By default, only distances strictly below threshold_mi are returned.
    Set include_non_overlaps=True to return all cross-utility pairs.
    Empty input or no qualifying pairs returns []. Input records are not modified.
    Invalid records raise ValueError with the project location.
    """
    threshold_mi = _threshold(threshold_mi)
    prepared, seen = [], set()
    for index, project in enumerate(projects):
        try:
            identifier, utility = project["project_id"], project["utility"]
            if not isinstance(identifier, str) or not identifier.strip():
                raise ValueError("project_id must be a nonempty string")
            if not isinstance(utility, str) or not utility.strip():
                raise ValueError("utility must be a nonempty string")
            if identifier in seen:
                raise ValueError(f"Duplicate project_id: {identifier}")
            seen.add(identifier)
            lat, lon = _point(project["lat_center"], project["lon_center"])
            prepared.append((project, lat, lon, _date(project.get("in_service_date"))))
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError(f"projects[{index}]: {exc}") from exc
    prepared.sort(key=lambda item: (item[0]["utility"], item[0]["project_id"]))
    results = []
    for (a, lat_a, lon_a, date_a), (b, lat_b, lon_b, date_b) in combinations(prepared, 2):
        if a["utility"] == b["utility"]:
            continue
        distance = calculate_distance(lat_a, lon_a, lat_b, lon_b, threshold_mi)
        if distance["is_overlap"] or include_non_overlaps:
            results.append({
                "project_id_a": a["project_id"], "project_id_b": b["project_id"],
                "utility_a": a["utility"], "utility_b": b["utility"],
                "project_name_a": a.get("project_name", ""),
                "project_name_b": b.get("project_name", ""),
                **distance,
                "time_gap_days": abs((date_a - date_b).days) if date_a and date_b else None,
            })
    return sorted(results, key=lambda pair: (pair["distance_mi"], pair["project_id_a"], pair["project_id_b"]))


def read_projects_from_excel(path: str | Path) -> list[Project]:
    """Read Projects_Overlaps.xlsx into dictionaries accepted by compare_projects.

    Requires openpyxl (python -m pip install openpyxl). Reads the projects sheet
    with project_id and utility columns. Both endpoint pairs produce their
    arithmetic midpoint; one endpoint becomes the center. If endpoints are absent,
    lat_center/lon_center must be numeric or have cached Excel formula values.
    Existing overlap results are ignored. The original file is never modified.
    Invalid rows raise ValueError with their Excel row number.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "projects" not in workbook.sheetnames:
            raise ValueError("Missing sheet: projects")
        rows = workbook["projects"].iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(rows, ())]
        if not {"project_id", "utility"} <= set(headers):
            raise ValueError("projects sheet requires project_id and utility columns")
        projects = []
        for row_number, values in enumerate(rows, 2):
            if all(value is None for value in values):
                continue
            row = dict(zip(headers, values))
            try:
                endpoints = []
                for suffix in ("a", "b"):
                    lat, lon = row.get(f"lat_{suffix}"), row.get(f"lon_{suffix}")
                    if lat is not None or lon is not None:
                        endpoints.append(_point(lat, lon))
                if endpoints:
                    lat, lon = (sum(coords) / len(endpoints) for coords in zip(*endpoints))
                else:
                    lat, lon = _point(row.get("lat_center"), row.get("lon_center"))
                for key in ("project_id", "utility"):
                    if not isinstance(row[key], str) or not row[key].strip():
                        raise ValueError(f"{key} must be a nonempty string")
                projects.append({
                    "project_id": row["project_id"].strip(), "utility": row["utility"].strip(),
                    "project_name": str(row.get("project_name") or ""),
                    "lat_center": lat, "lon_center": lon,
                    "in_service_date": _date(row.get("in_service_date")),
                })
            except (TypeError, ValueError) as exc:
                raise ValueError(f"projects row {row_number}: {exc}") from exc
        return projects
    finally:
        workbook.close()
